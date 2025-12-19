"""SOP-Generator mit Streamlit.

Dieses Skript erstellt eine Streamlit-App zum Generieren von
Standard Operating Procedures (SOPs) aus Word-Dokumenten.
"""

import io
from datetime import datetime

import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


def read_docx(file):
    """Liest ein Word-Dokument und gibt den Inhalt zurück.

    Args:
        file: Hochgeladene Datei (BytesIO)

    Returns:
        list: Liste von Absätzen (Paragraphs)
    """
    doc = Document(file)
    paragraphs = []

    for para in doc.paragraphs:
        if para.text.strip():  # Nur nicht-leere Absätze
            paragraphs.append({"text": para.text, "style": para.style.name})

    return paragraphs


def add_header(doc, metadata):
    """Fügt den SOP-Header hinzu.

    Args:
        doc: Document-Objekt
        metadata: Dictionary mit SOP-Metadaten
    """
    # Titel
    title = doc.add_heading(metadata["title"], level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Untertitel
    subtitle = doc.add_paragraph("Standard Operating Procedure")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(14)
    subtitle.runs[0].font.color.rgb = RGBColor(128, 128, 128)

    # Metadaten-Tabelle
    doc.add_paragraph()  # Leerzeile
    table = doc.add_table(rows=4, cols=2)
    table.style = "Light Grid Accent 1"

    # Tabelle füllen
    cells = [
        ("Version:", metadata["version"]),
        ("Autor:", metadata["author"]),
        ("Datum:", metadata["date"]),
        ("Status:", "Entwurf"),
    ]

    for i, (label, value) in enumerate(cells):
        table.rows[i].cells[0].text = label
        table.rows[i].cells[1].text = value
        # Erste Spalte fett
        table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph()  # Leerzeile


def add_content(doc, paragraphs):
    """Fügt den Hauptinhalt hinzu.

    Args:
        doc: Document-Objekt
        paragraphs: Liste von Absätzen
    """
    # Überschrift für Inhalt
    doc.add_heading("Inhalt", level=1)

    # Absätze hinzufügen
    for para in paragraphs:
        if "Heading" in para["style"]:
            # Überschrift
            level = int(para["style"].replace("Heading ", ""))
            doc.add_heading(para["text"], level=min(level, 3))
        else:
            # Normaler Absatz
            p = doc.add_paragraph(para["text"])
            p.style = "Normal"


def create_sop(content, metadata):
    """Erstellt ein formatiertes SOP-Dokument.

    Args:
        content: Liste von Absätzen
        metadata: Dictionary mit SOP-Metadaten

    Returns:
        BytesIO: Word-Dokument als Bytes
    """
    doc = Document()

    # Header hinzufügen
    add_header(doc, metadata)

    # Inhalt hinzufügen
    add_content(doc, content)

    # Dokument in BytesIO speichern
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    return buffer


# Seitenkonfiguration
st.set_page_config(
    page_title="SOP-Generator",
    page_icon="📄",
    layout="wide",
)

# Titel
st.title("📄 SOP-Generator")
st.markdown("Erstellen Sie formatierte Standard Operating Procedures aus Word-Dokumenten.")

# Zwei Spalten
col1, col2 = st.columns([1, 1])

with col1:
    st.subheader("1️⃣ Dokument hochladen")
    uploaded_file = st.file_uploader(
        "Wählen Sie ein Word-Dokument",
        type=["docx"],
        help="Laden Sie ein Word-Dokument mit dem Inhalt hoch",
    )

    if uploaded_file:
        st.success(f"✅ Datei hochgeladen: {uploaded_file.name}")

        # Dokument lesen
        content = read_docx(uploaded_file)
        st.info(f"📄 {len(content)} Absätze gefunden")

with col2:
    st.subheader("2️⃣ SOP-Informationen")

    with st.form("sop_form"):
        title = st.text_input("Titel", placeholder="z.B. Prozess XYZ")
        version = st.text_input("Version", value="1.0")
        author = st.text_input("Autor", placeholder="Ihr Name")
        date = st.date_input("Datum", value=datetime.now())

        submitted = st.form_submit_button("SOP erstellen", type="primary")

# SOP generieren
if uploaded_file and submitted:
    # Validierung
    if not title:
        st.error("❌ Bitte geben Sie einen Titel ein")
    elif not author:
        st.error("❌ Bitte geben Sie einen Autor ein")
    else:
        # Metadaten zusammenstellen
        metadata = {
            "title": title,
            "version": version,
            "author": author,
            "date": date.strftime("%d.%m.%Y"),
        }

        # SOP erstellen
        with st.spinner("SOP wird erstellt..."):
            sop_buffer = create_sop(content, metadata)

        st.success("✅ SOP erfolgreich erstellt!")

        # Download-Button
        st.download_button(
            label="📥 SOP herunterladen",
            data=sop_buffer,
            file_name=f"SOP_{title.replace(' ', '_')}_{version}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            type="primary",
        )

        # Vorschau
        with st.expander("👁️ Vorschau"):
            st.markdown(f"**Titel:** {title}")
            st.markdown(f"**Version:** {version}")
            st.markdown(f"**Autor:** {author}")
            st.markdown(f"**Datum:** {date.strftime('%d.%m.%Y')}")
            st.markdown("---")
            st.markdown(f"**Anzahl Absätze:** {len(content)}")

            # Erste 3 Absätze anzeigen
            st.markdown("**Erste Absätze:**")
            for i, para in enumerate(content[:3]):
                st.markdown(f"- {para['text'][:100]}...")

# Hilfe
st.markdown("---")
with st.expander("ℹ️ Hilfe"):
    st.markdown("""
    ### Wie verwende ich den SOP-Generator?

    1. **Dokument hochladen:** Laden Sie ein Word-Dokument (.docx) mit dem Inhalt hoch
    2. **Metadaten eingeben:** Füllen Sie Titel, Version, Autor und Datum aus
    3. **SOP erstellen:** Klicken Sie auf "SOP erstellen"
    4. **Download:** Laden Sie die formatierte SOP herunter

    ### Was macht der Generator?

    - Liest den Inhalt aus Ihrem Dokument
    - Fügt einen professionellen Header hinzu
    - Formatiert Überschriften und Absätze
    - Erstellt eine Metadaten-Tabelle
    - Generiert ein herunterladbares Word-Dokument

    ### Tipps

    - Verwenden Sie Überschriften in Ihrem Quelldokument für bessere Struktur
    - Der Generator behält die Überschriften-Hierarchie bei
    - Leere Absätze werden automatisch entfernt
    """)
